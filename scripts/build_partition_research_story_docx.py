"""Build the editable research narrative from its Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/research/partition-prefetch-research-story-zh.md'
DEST = ROOT / 'output/doc/glint_partition_prefetch_research_story_zh.docx'


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'PingFang TC')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (('Title', 23), ('Heading 1', 17), ('Heading 2', 13)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('17365D')
        doc.styles[name].paragraph_format.keep_with_next = True
    section.header.paragraphs[0].text = 'GLINT | Partition 供料架構與研究敘事'
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('研究工作稿 · ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    lines = SOURCE.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith('|'):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip())
                i += 1
            cells = [[c.strip() for c in r.strip('|').split('|')] for r in rows]
            cells = [r for r in cells if not all(re.fullmatch(r':?-+:?', c) for c in r)]
            table = doc.add_table(rows=0, cols=len(cells[0]))
            table.style = 'Light Shading Accent 1'
            for index, row in enumerate(cells):
                tr = table.add_row()
                for cell, value in zip(tr.cells, row):
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(9)
                if index == 0:
                    repeat = OxmlElement('w:tblHeader')
                    tr._tr.get_or_add_trPr().append(repeat)
                no_split = OxmlElement('w:cantSplit')
                tr._tr.get_or_add_trPr().append(no_split)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif line.startswith('!['):
            match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            path = (SOURCE.parent / match.group(2)).resolve()
            doc.add_picture(str(path), width=Cm(16.5))
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        else:
            doc.add_paragraph(line, 'Caption' if line.startswith('圖說：') else 'Normal')
    DEST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DEST)
    print(DEST)
    print(f'{len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables')


if __name__ == '__main__':
    build()
